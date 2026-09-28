import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

FILE_NAME = "vehicle_service.xlsx"


# ---------------- EXCEL SETUP ----------------

def create_excel_file():
    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Vehicle Service"

        headers = [
            "Vehicle No",
            "Owner Name",
            "Phone",
            "Vehicle Model",
            "Service Type",
            "Service Date",
            "Service Cost",
            "Status"
        ]

        ws.append(headers)
        wb.save(FILE_NAME)


create_excel_file()


# ---------------- MAIN WINDOW ----------------

root = tk.Tk()
root.title("Vehicle Service Management System")
root.geometry("1000x650")
root.resizable(False, False)


# ---------------- VARIABLES ----------------

vehicle_no = tk.StringVar()
owner_name = tk.StringVar()
phone = tk.StringVar()
vehicle_model = tk.StringVar()
service_type = tk.StringVar()
service_date = tk.StringVar()
service_cost = tk.StringVar()
status = tk.StringVar()


# ---------------- CLEAR FORM ----------------

def clear_form():
    vehicle_no.set("")
    owner_name.set("")
    phone.set("")
    vehicle_model.set("")
    service_type.set("")
    service_date.set("")
    service_cost.set("")
    status.set("")


# ---------------- ADD RECORD ----------------

def add_record():

    if (vehicle_no.get() == "" or
        owner_name.get() == "" or
        phone.get() == "" or
        vehicle_model.get() == "" or
        service_type.get() == "" or
        service_date.get() == "" or
        service_cost.get() == "" or
        status.get() == ""):

        messagebox.showwarning("Warning", "Please fill all fields.")
        return

    try:
        float(service_cost.get())
    except ValueError:
        messagebox.showerror("Error", "Service Cost must be a number.")
        return

    wb = load_workbook(FILE_NAME)
    ws = wb["Vehicle Service"]

    ws.append([
        vehicle_no.get(),
        owner_name.get(),
        phone.get(),
        vehicle_model.get(),
        service_type.get(),
        service_date.get(),
        service_cost.get(),
        status.get()
    ])

    wb.save(FILE_NAME)

    messagebox.showinfo("Success", "Vehicle service record added successfully.")

    clear_form()
    view_records()


# ---------------- VIEW RECORDS ----------------

def view_records():

    for item in tree.get_children():
        tree.delete(item)

    wb = load_workbook(FILE_NAME)
    ws = wb["Vehicle Service"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", tk.END, values=row)


# ---------------- SELECT RECORD ----------------

def select_record(event):

    selected = tree.focus()

    if not selected:
        return

    values = tree.item(selected, "values")

    if values:
        vehicle_no.set(values[0])
        owner_name.set(values[1])
        phone.set(values[2])
        vehicle_model.set(values[3])
        service_type.set(values[4])
        service_date.set(values[5])
        service_cost.set(values[6])
        status.set(values[7])


# ---------------- UPDATE RECORD ----------------

def update_record():

    selected = tree.focus()

    if not selected:
        messagebox.showwarning(
            "Warning",
            "Please select a record to update."
        )
        return

    values = tree.item(selected, "values")
    old_vehicle_no = values[0]

    wb = load_workbook(FILE_NAME)
    ws = wb["Vehicle Service"]

    for row in ws.iter_rows(min_row=2):

        if str(row[0].value) == str(old_vehicle_no):

            row[0].value = vehicle_no.get()
            row[1].value = owner_name.get()
            row[2].value = phone.get()
            row[3].value = vehicle_model.get()
            row[4].value = service_type.get()
            row[5].value = service_date.get()
            row[6].value = service_cost.get()
            row[7].value = status.get()

            break

    wb.save(FILE_NAME)

    messagebox.showinfo(
        "Success",
        "Record updated successfully."
    )

    clear_form()
    view_records()


# ---------------- DELETE RECORD ----------------

def delete_record():

    selected = tree.focus()

    if not selected:
        messagebox.showwarning(
            "Warning",
            "Please select a record to delete."
        )
        return

    values = tree.item(selected, "values")
    vehicle_number = values[0]

    confirm = messagebox.askyesno(
        "Confirm Delete",
        "Are you sure you want to delete this record?"
    )

    if not confirm:
        return

    wb = load_workbook(FILE_NAME)
    ws = wb["Vehicle Service"]

    for row in range(2, ws.max_row + 1):

        if str(ws.cell(row, 1).value) == str(vehicle_number):
            ws.delete_rows(row, 1)
            break

    wb.save(FILE_NAME)

    messagebox.showinfo(
        "Success",
        "Record deleted successfully."
    )

    clear_form()
    view_records()


# ---------------- SEARCH ----------------

def search_record():

    search_value = search_entry.get().lower()

    for item in tree.get_children():
        tree.delete(item)

    wb = load_workbook(FILE_NAME)
    ws = wb["Vehicle Service"]

    for row in ws.iter_rows(min_row=2, values_only=True):

        row_text = " ".join(
            str(value).lower() for value in row
        )

        if search_value in row_text:
            tree.insert("", tk.END, values=row)


# ---------------- LOGOUT ----------------

def logout():

    answer = messagebox.askyesno(
        "Logout",
        "Do you want to logout?"
    )

    if answer:
        dashboard_frame.pack_forget()
        login_frame.pack(fill="both", expand=True)

        username_entry.delete(0, tk.END)
        password_entry.delete(0, tk.END)


# ---------------- LOGIN ----------------

def login():

    username = username_entry.get()
    password = password_entry.get()

    if username == "admin" and password == "1234":

        login_frame.pack_forget()
        dashboard_frame.pack(fill="both", expand=True)

        view_records()

    else:
        messagebox.showerror(
            "Login Failed",
            "Invalid username or password."
        )


# ==================================================
# LOGIN PAGE
# ==================================================

login_frame = tk.Frame(root)
login_frame.pack(fill="both", expand=True)

login_title = tk.Label(
    login_frame,
    text="Vehicle Service Management System",
    font=("Arial", 24, "bold")
)
login_title.pack(pady=50)

login_box = tk.Frame(login_frame)
login_box.pack()

tk.Label(
    login_box,
    text="Username",
    font=("Arial", 12)
).grid(row=0, column=0, padx=10, pady=10)

username_entry = tk.Entry(
    login_box,
    width=30,
    font=("Arial", 12)
)
username_entry.grid(row=0, column=1, padx=10, pady=10)


tk.Label(
    login_box,
    text="Password",
    font=("Arial", 12)
).grid(row=1, column=0, padx=10, pady=10)

password_entry = tk.Entry(
    login_box,
    width=30,
    show="*",
    font=("Arial", 12)
)
password_entry.grid(row=1, column=1, padx=10, pady=10)


tk.Button(
    login_box,
    text="LOGIN",
    command=login,
    width=20,
    font=("Arial", 12, "bold")
).grid(row=2, column=0, columnspan=2, pady=20)


tk.Label(
    login_frame,
    text="Username: admin    Password: 1234",
    font=("Arial", 10)
).pack()


# ==================================================
# DASHBOARD
# ==================================================

dashboard_frame = tk.Frame(root)

# Header

header = tk.Frame(
    dashboard_frame,
    bd=2,
    relief=tk.RIDGE
)
header.pack(fill="x")

tk.Label(
    header,
    text="Vehicle Service Management System",
    font=("Arial", 20, "bold")
).pack(side="left", padx=20, pady=15)

tk.Button(
    header,
    text="Logout",
    command=logout,
    width=10
).pack(side="right", padx=20)


# ==================================================
# FORM
# ==================================================

form_frame = tk.LabelFrame(
    dashboard_frame,
    text="Vehicle Service Details",
    font=("Arial", 12, "bold")
)
form_frame.pack(fill="x", padx=10, pady=10)


tk.Label(form_frame, text="Vehicle No").grid(
    row=0, column=0, padx=10, pady=7
)

tk.Entry(
    form_frame,
    textvariable=vehicle_no,
    width=20
).grid(row=0, column=1, padx=10, pady=7)


tk.Label(form_frame, text="Owner Name").grid(
    row=0, column=2, padx=10, pady=7
)

tk.Entry(
    form_frame,
    textvariable=owner_name,
    width=20
).grid(row=0, column=3, padx=10, pady=7)


tk.Label(form_frame, text="Phone").grid(
    row=1, column=0, padx=10, pady=7
)

tk.Entry(
    form_frame,
    textvariable=phone,
    width=20
).grid(row=1, column=1, padx=10, pady=7)


tk.Label(form_frame, text="Vehicle Model").grid(
    row=1, column=2, padx=10, pady=7
)

tk.Entry(
    form_frame,
    textvariable=vehicle_model,
    width=20
).grid(row=1, column=3, padx=10, pady=7)


tk.Label(form_frame, text="Service Type").grid(
    row=2, column=0, padx=10, pady=7
)

service_combo = ttk.Combobox(
    form_frame,
    textvariable=service_type,
    values=[
        "General Service",
        "Oil Change",
        "Brake Service",
        "Engine Repair",
        "Tyre Service",
        "Battery Service"
    ],
    width=18,
    state="readonly"
)
service_combo.grid(row=2, column=1, padx=10, pady=7)


tk.Label(form_frame, text="Service Date").grid(
    row=2, column=2, padx=10, pady=7
)

tk.Entry(
    form_frame,
    textvariable=service_date,
    width=20
).grid(row=2, column=3, padx=10, pady=7)


tk.Label(form_frame, text="Service Cost").grid(
    row=3, column=0, padx=10, pady=7
)

tk.Entry(
    form_frame,
    textvariable=service_cost,
    width=20
).grid(row=3, column=1, padx=10, pady=7)


tk.Label(form_frame, text="Status").grid(
    row=3, column=2, padx=10, pady=7
)

status_combo = ttk.Combobox(
    form_frame,
    textvariable=status,
    values=[
        "Pending",
        "In Service",
        "Completed"
    ],
    width=18,
    state="readonly"
)
status_combo.grid(row=3, column=3, padx=10, pady=7)


# ==================================================
# BUTTONS
# ==================================================

button_frame = tk.Frame(dashboard_frame)
button_frame.pack(pady=5)

tk.Button(
    button_frame,
    text="Add Record",
    command=add_record,
    width=15
).grid(row=0, column=0, padx=5)

tk.Button(
    button_frame,
    text="Update",
    command=update_record,
    width=15
).grid(row=0, column=1, padx=5)

tk.Button(
    button_frame,
    text="Delete",
    command=delete_record,
    width=15
).grid(row=0, column=2, padx=5)

tk.Button(
    button_frame,
    text="Clear / Reset",
    command=clear_form,
    width=15
).grid(row=0, column=3, padx=5)

tk.Button(
    button_frame,
    text="View All",
    command=view_records,
    width=15
).grid(row=0, column=4, padx=5)


# ==================================================
# SEARCH
# ==================================================

search_frame = tk.Frame(dashboard_frame)
search_frame.pack(fill="x", padx=10, pady=5)

tk.Label(
    search_frame,
    text="Search:"
).pack(side="left")

search_entry = tk.Entry(
    search_frame,
    width=40
)
search_entry.pack(side="left", padx=10)

tk.Button(
    search_frame,
    text="Search",
    command=search_record
).pack(side="left")


# ==================================================
# TREEVIEW
# ==================================================

table_frame = tk.Frame(dashboard_frame)
table_frame.pack(fill="both", expand=True, padx=10, pady=5)

columns = (
    "Vehicle No",
    "Owner Name",
    "Phone",
    "Vehicle Model",
    "Service Type",
    "Service Date",
    "Service Cost",
    "Status"
)

tree = ttk.Treeview(
    table_frame,
    columns=columns,
    show="headings",
    height=12
)

for column in columns:

    tree.heading(
        column,
        text=column
    )

    tree.column(
        column,
        width=110
    )

tree.pack(
    side="left",
    fill="both",
    expand=True
)

scrollbar = ttk.Scrollbar(
    table_frame,
    orient="vertical",
    command=tree.yview
)

tree.configure(
    yscrollcommand=scrollbar.set
)

scrollbar.pack(
    side="right",
    fill="y"
)

tree.bind(
    "<ButtonRelease-1>",
    select_record
)


# ---------------- START ----------------

root.mainloop()